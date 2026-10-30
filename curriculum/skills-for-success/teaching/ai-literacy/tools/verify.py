"""Check real file types, media timing, named tasks and usable browser outputs."""
from pathlib import Path
from functools import partial
from http.server import SimpleHTTPRequestHandler,ThreadingHTTPServer
from threading import Thread
import json,re,subprocess,zipfile,hashlib
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from docx import Document
from pypdf import PdfReader
from playwright.sync_api import sync_playwright
R=Path(__file__).resolve().parents[1];B=R.parents[1];F=R/'fixtures';checks=[]
def check(name,condition,details=None):
 checks.append({'check':name,'passed':bool(condition),'details':details});assert condition,(name,details)
w=load_workbook(F/'event-budget.xlsx',data_only=False);s=w['Budget'];check('Workbook contains actual formulas',[s[f'D{i}'].value for i in [2,3,4,6,8]]==['=B2*C2','=B3*C3','=B4*C4','=SUM(D2:D4)','=D7-D6'])
check('Budget arithmetic and change cases',[24*6.5+24*3.25+38,24*6.5+26*3.25+38,26*6.5+26*3.25+38]==[272,278.5,291.5])
doc=Document(F/'event-brief.docx');check('DOCX opens with complete table',len(doc.tables)==1 and len(doc.tables[0].rows)==4)
check('PDF brief is readable and complete','Cedar Room' in PdfReader(F/'event-brief.pdf').pages[0].extract_text())
with zipfile.ZipFile(F/'ai-practice-pack.zip') as z:
 check('Practice ZIP CRC',z.testzip() is None);check('Practice ZIP contains actual media and formats',all(x in z.namelist() for x in ['event-brief.docx','event-budget.xlsx','file-formats-study.mp4','file-formats-study.mp3','file-formats-study.vtt']))
probe=json.loads(subprocess.check_output(['ffprobe','-v','error','-show_streams','-show_format','-of','json',str(F/'file-formats-study.mp4')]))
check('MP4 has audio and video',{x['codec_type'] for x in probe['streams']} >= {'audio','video'})
m=json.loads((F/'media-manifest.json').read_text());duration=float(probe['format']['duration']);check('Media duration matches storyboard',abs(duration-m['duration_seconds'])<.3,{'actual':duration,'planned':m['duration_seconds']})
cues=re.findall(r'(\d\d:\d\d:\d\d\.\d{3}) --> (\d\d:\d\d:\d\d\.\d{3})',(F/'file-formats-study.vtt').read_text());seconds=lambda t:sum(float(x)*n for x,n in zip(t.split(':'),[3600,60,1]));check('Captions ordered within media',bool(cues) and all(seconds(a)<seconds(b)<=duration+.1 for a,b in cues) and all(seconds(cues[i][1])<=seconds(cues[i+1][0])+.002 for i in range(len(cues)-1)))
for stem in ['learner-textbook','practitioner-guide']:
 soup=BeautifulSoup((B/(stem+'.html')).read_text(),'html.parser');text=soup.get_text(' ',strip=True)
 check(stem+' has no visible opaque task codes',not re.search(r'\b(?:CK-)?(?:RD|WR|NU|DG|PS|CM|CL|AD|CI|CP|AI)-[ABCR]?\d{2}\b',text))
 check(stem+' has sixteen named AI activities',all(soup.select(f'[data-activity-id="AI-{i:02}"]') for i in range(1,17)))
 check(stem+' has no repository/course dependency',not re.search(r'\brepository\b|\brepo\b|asphalt A1|Roadbuilding foundations',text,re.I))
 pdftext=' '.join(p.extract_text() or '' for p in PdfReader(B/(stem+'.pdf')).pages)
 check(stem+' PDF includes audio/video activities',all(x in pdftext for x in ['question-and-pause audio lesson','captioned visual explanation','Audit a study recording','Build a small study kit']))
class Quiet(SimpleHTTPRequestHandler):
 def log_message(self,*args):pass
server=ThreadingHTTPServer(('127.0.0.1',0),partial(Quiet,directory=str(B)));Thread(target=server.serve_forever,daemon=True).start()
base=f'http://127.0.0.1:{server.server_port}/'
try:
 with sync_playwright() as pw:
  browser=pw.chromium.launch(executable_path='/usr/bin/chromium',headless=True,args=['--no-sandbox']);page=browser.new_page()
  for width in [390,1280]:
   page.set_viewport_size({'width':width,'height':900})
   for path in ['learner/06-communication.html','learner/04-digital.html','practitioner/06-digital.html','teaching/ai-literacy/media.html']:
    page.goto(base+path,wait_until='networkidle');page.evaluate('document.fonts.ready');bad=page.evaluate('''() => ({overflow:document.documentElement.scrollWidth>innerWidth+1,badImages:[...document.images].filter(i=>i.complete&&i.naturalWidth===0).map(i=>i.src)})''');check(f'Browser {width} {path}',not bad['overflow'] and not bad['badImages'],bad)
  page.goto(base+'teaching/ai-literacy/media.html');page.locator('video').evaluate('v=>v.load()');page.wait_for_function('document.querySelector("video").readyState>=1');check('Browser decodes video metadata',page.locator('video').evaluate('v=>v.videoWidth===1280 && v.duration>50'))
  for p in sorted((B/'assets/diagrams').glob('*.svg')):
   if not p.stem.startswith(('ai-','quick-')):continue
   page.set_content(p.read_text());bad=page.evaluate('''() => {let s=document.querySelector('svg'),v=s.viewBox.baseVal;return [...s.querySelectorAll('text')].filter(t=>{let b=t.getBBox();return b.x<0||b.y<0||b.x+b.width>v.width||b.y+b.height>v.height}).map(t=>t.textContent)}''');check('SVG bounds '+p.name,not bad,bad)
  browser.close()
finally:server.shutdown();server.server_close()
(R/'verification.json').write_text(json.dumps({'status':'passed','checks':checks,'limits':'Structural, arithmetic, timing and browser checks; not an accessibility certification or a learner pilot.'},indent=2)+'\n');print(len(checks),'AI/media checks passed')
